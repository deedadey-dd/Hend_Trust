import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_callout_borders(cell, border_color="F59E0B"):
    """Sets a prominent left border and light/none for others."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def add_formatted_runs(paragraph, text, base_font_size=10, base_color=RGBColor(30, 41, 59), base_bold=False):
    """Parses markdown bold (**text**), italic (*text*), and inline code (`text`) into docx runs."""
    # Pattern to match **bold**, *italic*, and `code`
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.size = Pt(base_font_size)
        run.font.name = "Segoe UI"
        
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
            run.font.color.rgb = base_color
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
            run.font.color.rgb = base_color
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run.text = token[1:-1]
            run.bold = True
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size - 0.5)
            run.font.color.rgb = RGBColor(3, 99, 255) # Hendaxis blue
        else:
            run.text = token
            run.bold = base_bold
            run.font.color.rgb = base_color

def build_docx(md_filepath, output_docx_path):
    doc = docx.Document()

    # Configure Margins (0.8 inches)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = header_p.add_run("HENDAXIS TRUST  •  MASTER TERMS OF SERVICE & USER AGREEMENT")
        h_run.font.name = "Segoe UI"
        h_run.font.size = Pt(8.5)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(100, 116, 139) # Slate-500

        # Footer
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        f_run = footer_p.add_run("CONFIDENTIAL  •  FOR REGULATORY, LEGAL & OPERATIONAL REVIEW")
        f_run.font.name = "Segoe UI"
        f_run.font.size = Pt(8)
        f_run.font.color.rgb = RGBColor(148, 163, 184) # Slate-400

    # Colors
    c_primary = RGBColor(3, 99, 255)     # #0363ff Hendaxis Blue
    c_orange = RGBColor(255, 109, 29)    # #ff6d1d Hendaxis Orange
    c_dark = RGBColor(15, 23, 42)        # #0f172a Slate-900
    c_body = RGBColor(30, 41, 59)        # #1e293b Slate-800
    c_sub = RGBColor(71, 85, 105)        # #475569 Slate-600
    c_alert_text = RGBColor(146, 64, 14) # #92400e Amber-800

    with open(md_filepath, "r", encoding="utf-8") as f:
        md_text = f.read()

    lines = md_text.splitlines()
    in_alert = False
    alert_buffer = []

    for line in lines:
        stripped = line.strip()

        # Handle Alert Blockquotes
        if stripped.startswith("> [!IMPORTANT]") or stripped.startswith("> [!WARNING]") or stripped.startswith("> [!NOTE]"):
            in_alert = True
            alert_buffer = []
            continue

        if in_alert:
            if stripped.startswith(">"):
                clean_alert_line = stripped.lstrip(">").strip()
                if clean_alert_line:
                    alert_buffer.append(clean_alert_line)
                continue
            else:
                # Flush alert callout box
                in_alert = False
                if alert_buffer:
                    table = doc.add_table(rows=1, cols=1)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    table.autofit = False
                    cell = table.cell(0, 0)
                    cell.width = Inches(6.9)
                    set_cell_background(cell, "FFFBEB") # Light amber
                    set_callout_borders(cell, "F59E0B") # Amber-500
                    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)

                    # Heading inside callout
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(4)
                    c_head_run = cp.add_run("⚠️ CRITICAL LEGAL NOTICE & INDEMNITY REQUIREMENT")
                    c_head_run.font.name = "Segoe UI"
                    c_head_run.font.size = Pt(9.5)
                    c_head_run.font.bold = True
                    c_head_run.font.color.rgb = c_alert_text

                    for ab in alert_buffer:
                        p_call = cell.add_paragraph()
                        p_call.paragraph_format.space_before = Pt(2)
                        p_call.paragraph_format.space_after = Pt(2)
                        p_call.paragraph_format.line_spacing = 1.15
                        add_formatted_runs(p_call, ab, base_font_size=9, base_color=c_alert_text)

                    # Add empty spacer paragraph after table
                    p_space = doc.add_paragraph()
                    p_space.paragraph_format.space_before = Pt(0)
                    p_space.paragraph_format.space_after = Pt(4)
                    alert_buffer = []

        if not stripped:
            continue

        if stripped == "---":
            # Horizontal separator
            p_hr = doc.add_paragraph()
            p_hr.paragraph_format.space_before = Pt(6)
            p_hr.paragraph_format.space_after = Pt(6)
            hr_run = p_hr.add_run("_________________________________________________________________________________")
            hr_run.font.name = "Segoe UI"
            hr_run.font.size = Pt(8)
            hr_run.font.color.rgb = RGBColor(226, 232, 240)
            continue

        # Replace UI emojis with clean text tags
        cleaned = stripped
        cleaned = cleaned.replace("🛡️", "[Verified]")
        cleaned = cleaned.replace("🆕", "[New]")
        cleaned = cleaned.replace("📜", "")
        cleaned = cleaned.replace("⚖️", "[Legal]")

        # Document Title (# Title)
        if stripped.startswith("# "):
            title_text = cleaned.lstrip("# ").strip()
            
            p_super = doc.add_paragraph()
            p_super.paragraph_format.space_before = Pt(0)
            p_super.paragraph_format.space_after = Pt(2)
            s_run = p_super.add_run("HENDAXIS TRUST")
            s_run.font.name = "Segoe UI"
            s_run.font.size = Pt(11)
            s_run.font.bold = True
            s_run.font.color.rgb = c_orange

            p_title = doc.add_paragraph()
            p_title.paragraph_format.space_before = Pt(0)
            p_title.paragraph_format.space_after = Pt(6)
            t_run = p_title.add_run(title_text)
            t_run.font.name = "Segoe UI"
            t_run.font.size = Pt(20)
            t_run.font.bold = True
            t_run.font.color.rgb = c_primary
            continue

        # Section Heading (## Section)
        if stripped.startswith("## "):
            h1_text = cleaned.lstrip("# ").strip()
            p_h1 = doc.add_paragraph()
            p_h1.paragraph_format.space_before = Pt(14)
            p_h1.paragraph_format.space_after = Pt(4)
            p_h1.paragraph_format.keep_with_next = True
            h1_run = p_h1.add_run(h1_text)
            h1_run.font.name = "Segoe UI"
            h1_run.font.size = Pt(13)
            h1_run.font.bold = True
            h1_run.font.color.rgb = c_dark
            continue

        # Subsection Heading (### Subsection)
        if stripped.startswith("### "):
            h2_text = cleaned.lstrip("# ").strip()
            p_h2 = doc.add_paragraph()
            p_h2.paragraph_format.space_before = Pt(10)
            p_h2.paragraph_format.space_after = Pt(3)
            p_h2.paragraph_format.keep_with_next = True
            h2_run = p_h2.add_run(h2_text)
            h2_run.font.name = "Segoe UI"
            h2_run.font.size = Pt(10.5)
            h2_run.font.bold = True
            h2_run.font.color.rgb = c_primary
            continue

        # Subtitle / Metadata (Last Updated / Effective Date)
        if "Last Updated:" in stripped or "Effective Date:" in stripped:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(10)
            add_formatted_runs(p_sub, cleaned, base_font_size=9.5, base_color=c_sub)
            continue

        # Bullet List Items
        if stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = cleaned[2:].strip()
            p_bullet = doc.add_paragraph(style='List Bullet')
            p_bullet.paragraph_format.space_before = Pt(2)
            p_bullet.paragraph_format.space_after = Pt(3)
            p_bullet.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p_bullet, bullet_text, base_font_size=9.5, base_color=c_body)
            continue

        # Numbered List Items (e.g. "1. Item")
        num_match = re.match(r'^(\d+\.)\s*(.*)', cleaned)
        if num_match:
            num_prefix = num_match.group(1)
            num_body = num_match.group(2)
            p_num = doc.add_paragraph()
            p_num.paragraph_format.left_indent = Inches(0.25)
            p_num.paragraph_format.space_before = Pt(2)
            p_num.paragraph_format.space_after = Pt(3)
            p_num.paragraph_format.line_spacing = 1.15
            
            p_num_run = p_num.add_run(f"{num_prefix} ")
            p_num_run.font.name = "Segoe UI"
            p_num_run.font.size = Pt(9.5)
            p_num_run.font.bold = True
            p_num_run.font.color.rgb = c_dark
            
            add_formatted_runs(p_num, num_body, base_font_size=9.5, base_color=c_body)
            continue

        # Standard Paragraph Body
        p_body = doc.add_paragraph()
        p_body.paragraph_format.space_before = Pt(3)
        p_body.paragraph_format.space_after = Pt(5)
        p_body.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p_body, cleaned, base_font_size=9.5, base_color=c_body)

    # Save Word Document
    doc.save(output_docx_path)
    print(f"Successfully generated DOCX at: {output_docx_path}")

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    md_path = os.path.join(base_dir, "docs", "TERMS_AND_CONDITIONS.md")
    docx_path = os.path.join(base_dir, "docs", "HENDAXIS_TRUST_TERMS_AND_CONDITIONS.docx")
    build_docx(md_path, docx_path)
